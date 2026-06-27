from __future__ import annotations

from typing import Any

from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE


def get_shape(shape: Any) -> str:
    """判断 PPTX shape 的粗粒度类型。

    Args:
        shape: python-pptx 的 shape 对象。

    Returns:
        shape 类型字符串。可能值包括 `text`、`table`、`chart-bar`、
        `chart-line`、`chart-pie`、`other`。
    """
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        return "table"

    if shape.shape_type == MSO_SHAPE_TYPE.CHART and hasattr(shape, "chart"):
        chart_type = shape.chart.chart_type
        if chart_type in _chart_type_values(
            "COLUMN_CLUSTERED",
            "COLUMN_STACKED",
            "COLUMN_STACKED_100",
            "BAR_CLUSTERED",
            "BAR_STACKED",
            "BAR_STACKED_100",
        ):
            return "chart-bar"
        if chart_type in _chart_type_values(
            "LINE",
            "LINE_MARKERS",
            "LINE_STACKED",
            "LINE_STACKED_100",
            "LINE_MARKERS_STACKED",
            "LINE_MARKERS_STACKED_100",
        ):
            return "chart-line"
        if chart_type in _chart_type_values("PIE", "PIE_EXPLODED", "DOUGHNUT"):
            return "chart-pie"
        return "other"

    if shape.has_text_frame and shape.text.strip():
        return "text"

    return "other"


def extract_body_table(shape: Any, role: str) -> tuple[list[str], list[list[Any]]]:
    """从 table 或 chart shape 中抽取二维数据。

    Args:
        shape: python-pptx 的 table/chart shape 对象。
        role: body 元素 role，用于区分 table 与 chart。

    Returns:
        二元组 `(header, rows)`。`header` 是 CSV 表头，`rows` 是数据行。
    """
    if role == "table":
        raw_rows = [
            [
                str(shape.table.cell(row_idx, col_idx).text).strip()
                for col_idx in range(len(shape.table.columns))
            ]
            for row_idx in range(len(shape.table.rows))
        ]
        if not raw_rows or not raw_rows[0]:
            raise ValueError("Parser cannot export an empty PPT table.")
        return raw_rows[0], raw_rows[1:]

    chart = shape.chart
    categories = [str(category.label) for category in chart.plots[0].categories]
    series = list(chart.series)
    if not categories or not series:
        raise ValueError("Parser cannot export an empty chart.")
    header = ["category"] + [str(item.name) for item in series]
    rows = []
    for idx, category in enumerate(categories):
        rows.append(
            [category]
            + [item.values[idx] if idx < len(item.values) else "" for item in series]
        )
    return header, rows


def _chart_type_values(*names: str) -> set[Any]:
    values = set()
    for name in names:
        if hasattr(XL_CHART_TYPE, name):
            values.add(getattr(XL_CHART_TYPE, name))
    return values
